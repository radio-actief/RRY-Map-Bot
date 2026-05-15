#!/usr/bin/env python3
"""
Generate RRY-Map-Bot/data/be-locode.json: exactly 565 Belgian municipalities.

Source of truth for the set: StatBel NIS6 sheet (CNIS5_2025 parents).
Province per gemeente: REFNIS_2025_CITIES.xlsx (sheet walk).
UN/LOCODE code-list.csv: 3-letter Location + default coordinates where available.
Manual tables: be_locode_nis5_un_overrides.py + legacy LOCODE_* in this file.
Optional: ``--osm-pbf`` uses OpenStreetMap admin_level=8 + ``ref:INS`` polygons (``osmium``;
see ``requirements-dev.txt``) for map pins: **settlement** ``place=*`` nodes inside each polygon
when available, else polygon **representative_point**, else UN/LOCODE.
Optional ``--sectors-geojson`` (StatBel statistical sectors, EPSG:3812) supplies an
area-weighted centroid per gemeente to break ties between multiple settlements.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import unicodedata
import urllib.request
from collections import defaultdict
from pathlib import Path

try:
    import openpyxl
except ImportError as e:  # pragma: no cover
    raise SystemExit(
        "openpyxl is required. Install with: pip install openpyxl\n"
        "Or use the project venv: python3 -m venv .venv && .venv/bin/pip install -r requirements.txt"
    ) from e

from be_locode_nis5_un_overrides import (
    NIS5_TO_UN_LOCATION,
    NIS5_UN_DISAMBIG,
)

SCRIPT_DIR = Path(__file__).resolve().parent
DATA_DIR = SCRIPT_DIR.parent / "data"
INPUTS_DIR = DATA_DIR / "inputs"


def _default_input_path(filename: str) -> Path:
    """Prefer ``data/inputs/<file>`` when present, else ``data/statbel/<file>``."""
    p_in = INPUTS_DIR / filename
    if p_in.is_file():
        return p_in
    return DATA_DIR / "statbel" / filename


DEFAULT_NIS6 = _default_input_path("NIS6withnamefrom01012025.xlsx")
DEFAULT_REFNIS = _default_input_path("REFNIS_2025_CITIES.xlsx")
UN_LOCODE_CSV = "https://raw.githubusercontent.com/datasets/un-locode/master/data/code-list.csv"

OUTPUT = DATA_DIR / "be-locode.json"
DEFAULT_SECTORS_GEOJSON = _default_input_path(
    "sh_statbel_statistical_sectors_3812_20250101.geojson"
)

FLEMISH_PROVINCES = frozenset({"be-van", "be-vbr", "be-vov", "be-vwv", "be-vli"})
WALLOON_PROVINCES = frozenset({"be-wbr", "be-wht", "be-wlg", "be-wna", "be-wlx"})

# Legacy UN Location (3 letters) coordinate fixes (still keyed by UN Location)
LOCODE_COORD_OVERRIDES = {
    "TNO": (49.717, 5.417),
    "VIT": (50.280, 6.127),
    "KZN": (50.944, 5.189),
    "PUR": (50.321, 5.262),
    "SDE": (50.472, 4.797),
}

PROVINCE_NL_TO_MESH = {
    "Provincie Antwerpen": "be-van",
    "Provincie Vlaams-Brabant": "be-vbr",
    "Provincie West-Vlaanderen": "be-vwv",
    "Provincie Oost-Vlaanderen": "be-vov",
    "Provincie Limburg": "be-vli",
    "Provincie Waals-Brabant": "be-wbr",
    "Provincie Henegouwen": "be-wht",
    "Provincie Luik": "be-wlg",
    "Provincie Luxemburg": "be-wlx",
    "Provincie Namen": "be-wna",
}


def norm_key(s: str) -> str:
    s = (s or "").lower().strip()
    s = "".join(
        c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn"
    )
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def parse_unlocode_coords(s):
    if not s or not str(s).strip():
        return None, None
    s = str(s).strip()
    parts = s.split()
    if len(parts) != 2:
        return None, None
    try:
        lat_s, lon_s = parts[0], parts[1]
        lat_val = int(lat_s[:4])
        lat_dir = lat_s[-1]
        lon_val = int(lon_s[:5])
        lon_dir = lon_s[-1]
        lat = (lat_val // 100) + (lat_val % 100) / 60
        lon = (lon_val // 100) + (lon_val % 100) / 60
        if lat_dir == "S":
            lat = -lat
        if lon_dir == "W":
            lon = -lon
        return lat, lon
    except (ValueError, IndexError):
        return None, None


def coords_for_un_location(loc_code: str, coords_str: str):
    loc_u = loc_code.strip().upper()
    if loc_u in LOCODE_COORD_OVERRIDES:
        lat, lon = LOCODE_COORD_OVERRIDES[loc_u]
        return round(lat, 5), round(lon, 5)
    lat, lon = parse_unlocode_coords(coords_str)
    if lat is None:
        return None, None
    return round(lat, 5), round(lon, 5)


def norm_nis_cell(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v == int(v):
        v = int(v)
    s = str(v).strip()
    if s.endswith(".0") and s[:-2].isdigit():
        s = s[:-2]
    return s


def refnis_nis5_to_province(refnis_path: Path) -> dict[str, str]:
    wb = openpyxl.load_workbook(refnis_path, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    current: str | None = None
    out: dict[str, str] = {}
    for r in rows[1:]:
        if not r:
            continue
        nis = norm_nis_cell(r[3])
        nl = (r[4] or "").strip()
        fr = (r[1] or "").strip()
        if nl in PROVINCE_NL_TO_MESH:
            current = PROVINCE_NL_TO_MESH[nl]
            continue
        if nl == "BRUSSELS HOOFDSTEDELIJK GEWEST" or fr == "RÉGION DE BRUXELLES-CAPITALE":
            current = "be-bru"
            continue
        if nis == "21000" and "Arrondissement" in nl:
            current = "be-bru"
            continue
        if "Arrondissement" in nl or "Arrondissement" in fr:
            continue
        if nis and len(nis) == 5 and nis.isdigit():
            if int(nis) < 11001:
                continue
            if nis.endswith("000") and nis not in ("01000", "04000", "02000"):
                continue
            if current is None:
                continue
            out[nis] = current
    return out


def load_nis6_groups(nis6_path: Path) -> dict[str, list[tuple]]:
    """CNIS5 -> list of rows (T_MUN_NL, T_MUN_FR, T_MUN_DE, T_NIS6_NL, T_NIS6_FR)."""
    wb = openpyxl.load_workbook(nis6_path, read_only=True)
    ws = wb.active
    groups: dict[str, list] = defaultdict(list)
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if not row or row[0] is None:
            continue
        nis5 = norm_nis_cell(row[0])
        if not nis5 or len(nis5) != 5 or not nis5.isdigit():
            continue
        groups[nis5].append(
            (
                (row[1] or "").strip(),
                (row[2] or "").strip(),
                (row[3] or "").strip(),
                (row[5] or "").strip(),
                (row[6] or "").strip(),
            )
        )
    return dict(groups)


def un_name_tokens(name: str) -> set[str]:
    """Normalized tokens from UN Name for matching Dutch/French gemeente labels."""
    name = name or ""
    toks: set[str] = set()
    toks.add(norm_key(name))
    for part in re.split(r"[/]", name):
        part = part.strip()
        if not part:
            continue
        sub = re.split(r"\(", part, maxsplit=1)[0].strip()
        if sub:
            toks.add(norm_key(sub))
        inner = re.findall(r"\(([^)]+)\)", part)
        for inn in inner:
            inn = inn.strip()
            if inn:
                toks.add(norm_key(inn))
    return {t for t in toks if t}


def fetch_un_belgium_rows():
    with urllib.request.urlopen(UN_LOCODE_CSV, timeout=60) as resp:
        reader = csv.DictReader(line.decode("utf-8") for line in resp)
        return [r for r in reader if r.get("Country") == "BE" and r.get("Location")]


def gemeente_match_keys(nl: str, fr: str) -> set[str]:
    keys = {norm_key(nl), norm_key(fr)} - {""}
    for s in (nl, fr):
        for part in re.split(r"[-–]", s or ""):
            k = norm_key(part.strip())
            if len(k) >= 3:
                keys.add(k)
    return keys


def pick_un_for_gemeente(
    be_rows: list[dict],
    nis5: str,
    nl: str,
    fr: str,
    tokens_index: dict[str, list[dict]],
) -> tuple[str, dict] | tuple[None, None]:
    if nis5 in NIS5_TO_UN_LOCATION:
        loc = NIS5_TO_UN_LOCATION[nis5].upper()
        for r in be_rows:
            if r.get("Location", "").strip().upper() == loc:
                return loc, r
        return loc, {
            "Location": loc,
            "Name": nl,
            "Subdivision": "",
            "Coordinates": "",
            "Function": "23-----",
        }

    if nis5 in NIS5_UN_DISAMBIG:
        loc = NIS5_UN_DISAMBIG[nis5].upper()
        for r in be_rows:
            if r.get("Location", "").strip().upper() == loc:
                return loc, r

    keys = gemeente_match_keys(nl, fr)
    candidates: list[dict] = []
    seen_loc: set[str] = set()
    for k in keys:
        for r in tokens_index.get(k, []):
            loc = r.get("Location", "").strip().upper()
            if len(loc) != 3 or loc in seen_loc:
                continue
            seen_loc.add(loc)
            candidates.append(r)

    if not candidates:
        return None, None

    locs = {r["Location"].strip().upper() for r in candidates}
    if len(locs) == 1:
        loc = next(iter(locs))
        for r in candidates:
            if r["Location"].strip().upper() == loc:
                return loc, r

    if nis5 in NIS5_UN_DISAMBIG:
        loc = NIS5_UN_DISAMBIG[nis5].upper()
        for r in candidates:
            if r["Location"].strip().upper() == loc:
                return loc, r

    def score_row(r: dict) -> tuple:
        func = r.get("Function") or ""
        road = 1 if ("2" in func or "3" in func) else 0
        not_port_only = 0 if (func.startswith("1") and "3" not in func and "2" not in func) else 1
        status = r.get("Status") or ""
        st = 1 if status in ("AI", "RN", "RL") else 0
        return (road, not_port_only, st)

    best = max(candidates, key=score_row)
    return best["Location"].strip().upper(), best


def primary_municipality_label(province_code: str, nl_mun: str, fr_mun: str) -> str:
    """Official display string for ``plaats`` / ``gemeente`` (NL in Flanders + Brussels, FR in Wallonia)."""
    nl = (nl_mun or "").strip()
    fr = (fr_mun or "").strip()
    if province_code in FLEMISH_PROVINCES:
        return nl or fr
    if province_code in WALLOON_PROVINCES:
        return fr or nl
    if province_code == "be-bru":
        return nl or fr
    return nl or fr


def build_tokens_index(be_rows: list[dict]) -> dict[str, list[dict]]:
    idx: dict[str, list[dict]] = defaultdict(list)
    for r in be_rows:
        loc = r.get("Location", "").strip()
        if len(loc) != 3:
            continue
        func = r.get("Function", "")
        if func and func.startswith("1") and "3" not in func and "2" not in func:
            continue
        for tok in un_name_tokens(r.get("Name", "")):
            if len(tok) >= 3:
                idx[tok].append(r)
    return idx


# NIS6 sheet sometimes carries administrative noise; not useful for search.
_SKIP_ALIAS_NORM = {
    norm_key("Modification de limite communale"),
    norm_key("MODIFICATION DE LIMITE COMMUNALE"),
}


def build_aliases(
    _nis5: str,
    rows: list[tuple],
    primary: str,
    nl_mun: str,
    fr_mun: str,
) -> list[str]:
    seen: set[str] = set()
    order: list[str] = []

    def add(s: str):
        s = (s or "").strip().rstrip("*").strip()
        if len(s) < 2:
            return
        key = norm_key(s)
        if key in seen or key in _SKIP_ALIAS_NORM:
            return
        if "limite communale" in key or "grenswijziging" in key:
            return
        seen.add(key)
        order.append(s)

    for mun_nl, mun_fr, mun_de, sec_nl, sec_fr in rows:
        add(sec_nl)
        add(sec_fr)
        if mun_de and norm_key(mun_de) not in {norm_key(nl_mun), norm_key(fr_mun)}:
            add(mun_de)
    if nl_mun and norm_key(nl_mun) != norm_key(primary):
        add(nl_mun)
    if fr_mun and norm_key(fr_mun) != norm_key(primary):
        add(fr_mun)
    return [x for x in order if norm_key(x) != norm_key(primary)]


def main():
    ap = argparse.ArgumentParser(description="Generate be-locode.json (565 gemeenten).")
    ap.add_argument(
        "--nis6",
        type=Path,
        default=DEFAULT_NIS6,
        help="NIS6 xlsx (default: data/inputs/ then data/statbel/)",
    )
    ap.add_argument(
        "--refnis",
        type=Path,
        default=DEFAULT_REFNIS,
        help="REFNIS xlsx (default: data/inputs/ then data/statbel/)",
    )
    ap.add_argument("-o", "--output", type=Path, default=OUTPUT)
    ap.add_argument(
        "--osm-pbf",
        type=Path,
        default=None,
        help="Belgium .osm.pbf: settlement place=* inside admin polygon, else representative_point, "
        "else UN. Requires: pip install -r requirements-dev.txt",
    )
    ap.add_argument(
        "--osm-prefer",
        choices=("auto", "ogr", "osmium"),
        default="auto",
        help="With --osm-pbf: geometry extraction backend (default: auto)",
    )
    ap.add_argument(
        "--sectors-geojson",
        type=Path,
        default=None,
        help="StatBel statistical sectors GeoJSON (EPSG:3812) for settlement tie-break centroids "
        "(default: same basename under data/inputs/ then data/statbel/ if present; use with --osm-pbf only)",
    )
    ap.add_argument(
        "--no-sectors",
        action="store_true",
        help="Do not load the default sectors GeoJSON for tie-break (even if it exists)",
    )
    args = ap.parse_args()

    sectors_path = args.sectors_geojson
    if args.no_sectors:
        sectors_path = None
    elif sectors_path is None and DEFAULT_SECTORS_GEOJSON.is_file():
        sectors_path = DEFAULT_SECTORS_GEOJSON

    if not args.nis6.is_file():
        raise SystemExit(f"Missing NIS6 file: {args.nis6}")
    if not args.refnis.is_file():
        raise SystemExit(f"Missing REFNIS file: {args.refnis}")

    osm_geoms: dict | None = None
    if args.osm_pbf:
        if not args.osm_pbf.is_file():
            raise SystemExit(f"--osm-pbf not found: {args.osm_pbf}")
        try:
            from osm_pbf_gemeenten import load_nis5_geometries_with_meta
        except ImportError as e:
            raise SystemExit(
                "OSM PBF requires the `osmium` package (not in default API image). "
                "Install: pip install -r requirements-dev.txt\n"
                f"Detail: {e}"
            ) from e
        osm_geoms, osm_meta = load_nis5_geometries_with_meta(
            args.osm_pbf, prefer=args.osm_prefer
        )
        print(
            f"Loaded OSM admin_level=8 polygons: backend={osm_meta.get('backend')!r}, "
            f"{len(osm_geoms)} by ref:INS"
        )
        if osm_meta.get("ogr_error"):
            print(f"  (ogr2ogr not used: {osm_meta['ogr_error'][:120]})")

    sector_centroids: dict[str, tuple[float, float]] | None = None
    if args.osm_pbf and sectors_path is not None:
        if not sectors_path.is_file():
            raise SystemExit(f"--sectors-geojson not found: {sectors_path}")
        from statbel_sector_centroids import load_nis5_area_weighted_centroids

        sector_centroids = load_nis5_area_weighted_centroids(sectors_path)
        print(f"Loaded StatBel sector tie-break centroids: {len(sector_centroids)} gemeenten")

    settlement_coords: dict[str, tuple[float, float]] = {}
    if args.osm_pbf and osm_geoms is not None:
        from osm_pbf_settlements import load_nis5_settlement_coords

        settlement_coords, _set_meta, set_notes = load_nis5_settlement_coords(
            args.osm_pbf,
            osm_geoms,
            weighted_centroids_4326=sector_centroids,
        )
        print(f"Resolved OSM settlement coords: {len(settlement_coords)} gemeenten")
        for line in set_notes:
            print(f"  {line}")

    nis5_to_prov = refnis_nis5_to_province(args.refnis)
    groups = load_nis6_groups(args.nis6)
    if len(groups) != 565:
        raise SystemExit(f"Expected 565 CNIS5 parents, got {len(groups)}")

    print("Fetching UN/LOCODE Belgium rows...")
    be_rows = fetch_un_belgium_rows()
    tokens_index = build_tokens_index(be_rows)

    locations: list[dict] = []
    missing_un: list[str] = []
    coord_from_settlement = 0
    coord_from_osm_rep = 0
    coord_from_un_only = 0

    for nis5 in sorted(groups.keys()):
        rows = groups[nis5]
        nl_mun = (rows[0][0] or "").strip()
        fr_mun = (rows[0][1] or "").strip()
        prov = nis5_to_prov.get(nis5)
        if not prov:
            raise SystemExit(f"No province for NIS5 {nis5} {nl_mun}")

        loc_code, un_row = pick_un_for_gemeente(be_rows, nis5, nl_mun, fr_mun, tokens_index)
        if not loc_code:
            missing_un.append(f"{nis5} {nl_mun} / {fr_mun}")
            continue

        coords_str = (un_row or {}).get("Coordinates") or ""
        lat, lon = coords_for_un_location(loc_code, coords_str)

        if osm_geoms is not None:
            if nis5 in settlement_coords:
                slat, slon = settlement_coords[nis5]
                lat, lon = round(float(slat), 5), round(float(slon), 5)
                coord_from_settlement += 1
            else:
                g = osm_geoms.get(nis5)
                if g is not None and not g.is_empty:
                    try:
                        rp = g.representative_point()
                        lat, lon = round(float(rp.y), 5), round(float(rp.x), 5)
                        coord_from_osm_rep += 1
                    except Exception:
                        coord_from_un_only += 1
                else:
                    coord_from_un_only += 1

        province_code = prov
        display = primary_municipality_label(province_code, nl_mun, fr_mun)

        city_code = "be" + loc_code.lower()
        aliases = build_aliases(nis5, rows, display, nl_mun, fr_mun)

        rec = {
            "nis5": nis5,
            "plaats": display,
            "gemeente": display,
            "province_code": province_code,
            "city_code": city_code,
            "search_aliases": aliases,
        }
        if lat is not None and lon is not None:
            rec["lat"] = lat
            rec["lon"] = lon
        locations.append(rec)

    if missing_un:
        print("FAILED — no UN/synthetic row for:")
        for m in missing_un:
            print(" ", m)
        raise SystemExit(1)

    locations.sort(key=lambda x: (norm_key(x["plaats"]), x["city_code"]))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with open(args.output, "w", encoding="utf-8") as f:
        json.dump(locations, f, ensure_ascii=False, indent=0)

    print(f"Wrote {len(locations)} municipalities to {args.output}")
    if osm_geoms is not None:
        print(
            "Coordinate sources: "
            f"{coord_from_settlement} OSM settlement (place=*), "
            f"{coord_from_osm_rep} OSM polygon representative_point, "
            f"{coord_from_un_only} UN/LOCODE fallback (no OSM polygon or geom error)"
        )


if __name__ == "__main__":
    main()
